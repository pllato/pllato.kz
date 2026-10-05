from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "Pllato_Schet_PNK_Advance_250000.docx"
APP_OUT = ROOT / "app" / "schet-pnk-advance.docx"

NAVY = "17253D"
PALE = "F3F6FA"
BORDER = "D9D9D9"
MUTED = RGBColor(94, 108, 128)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def font(run, size=10.5, bold=False, color=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color) if isinstance(color, str) else color
    return run


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "nil")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(20, 24, 32)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.05

title_style = doc.styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

# Header
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
font(hp.add_run("ELC ALMATY"), 9.5, True, NAVY)
font(hp.add_run("   разработка систем управления"), 9, False, MUTED)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
p.add_run("Счет на предоплату")
remove_paragraph_border(p)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
font(p.add_run("№ PNK-С-2026-01 от 5 октября 2026 года"), 12, True, NAVY)

meta = doc.add_table(rows=2, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
set_table_borders(meta)
for row in meta.rows:
    row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_width(row.cells[0], 1.45)
    set_width(row.cells[1], 5.45)
    for c in row.cells:
        set_cell_margins(c)
for i, (label, value) in enumerate([
    ("Поставщик", "ТОО «ELC ALMATY», БИН 151040007861"),
    ("Покупатель", "Общественное объединение «Палата Налоговых Консультантов» (республиканский статус), БИН 020740002798"),
]):
    set_cell_shading(meta.cell(i, 0), PALE)
    p0 = meta.cell(i, 0).paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font(p0.add_run(label), 10, True, NAVY)
    p1 = meta.cell(i, 1).paragraphs[0]
    font(p1.add_run(value), 10, False)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(5)
font(p.add_run("Основание оплаты"), 11, True)
font(p.add_run("  Договор на разработку программного обеспечения № PNK-2026-01 от 5 октября 2026 года"), 10.5)

items = doc.add_table(rows=2, cols=5)
items.alignment = WD_TABLE_ALIGNMENT.CENTER
items.autofit = False
set_table_borders(items)
widths = [0.45, 3.65, 0.65, 1.0, 1.15]
for row in items.rows:
    for cell, w in zip(row.cells, widths):
        set_width(cell, w)
        set_cell_margins(cell, top=125, bottom=125)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
headers = ["№", "Наименование", "Кол во", "Цена", "Сумма"]
for idx, text in enumerate(headers):
    c = items.cell(0, idx)
    set_cell_shading(c, NAVY)
    p0 = c.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p0.add_run(text), 9.5, True, "FFFFFF")
set_repeat_header(items.rows[0])
row = items.rows[1]
values = [
    "1",
    "Предоплата 10 процентов за разработку портала управления консалтингом и реестром членов Палаты",
    "1 услуга",
    "250 000 ₸",
    "250 000 ₸",
]
for idx, text in enumerate(values):
    p0 = row.cells[idx].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 1 else WD_ALIGN_PARAGRAPH.CENTER
    font(p0.add_run(text), 9.7, idx in (4,))

totals = doc.add_table(rows=3, cols=2)
totals.alignment = WD_TABLE_ALIGNMENT.RIGHT
totals.autofit = False
set_table_borders(totals)
for row in totals.rows:
    set_width(row.cells[0], 1.65)
    set_width(row.cells[1], 1.65)
    for c in row.cells:
        set_cell_margins(c, top=85, bottom=85)
labels = [("Итого", "250 000 ₸"), ("НДС", "Без НДС"), ("К оплате", "250 000 ₸")]
for idx, (label, value) in enumerate(labels):
    if idx == 2:
        set_cell_shading(totals.cell(idx, 0), NAVY)
        set_cell_shading(totals.cell(idx, 1), NAVY)
        color = "FFFFFF"
    else:
        set_cell_shading(totals.cell(idx, 0), PALE)
        color = NAVY
    p0 = totals.cell(idx, 0).paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(p0.add_run(label), 10, True, color)
    p1 = totals.cell(idx, 1).paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(p1.add_run(value), 10, True, color if idx == 2 else None)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(8)
font(p.add_run("Всего к оплате 250 000 двести пятьдесят тысяч тенге 00 тиын. Без НДС."), 10.5, True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
font(p.add_run("Банковские реквизиты для оплаты"), 12, True)

bank = doc.add_table(rows=5, cols=2)
bank.alignment = WD_TABLE_ALIGNMENT.CENTER
bank.autofit = False
set_table_borders(bank)
bank_rows = [
    ("Получатель", "ТОО «ELC ALMATY», БИН 151040007861"),
    ("ИИК", "KZ82722S000033379752"),
    ("Банк", "АО «Kaspi Bank»"),
    ("БИК и КБе", "CASPKZKA, КБе 17"),
    ("КНП", "851"),
]
for idx, (label, value) in enumerate(bank_rows):
    set_width(bank.cell(idx, 0), 1.45)
    set_width(bank.cell(idx, 1), 5.45)
    for c in bank.rows[idx].cells:
        set_cell_margins(c, top=75, bottom=75)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(bank.cell(idx, 0), PALE)
    font(bank.cell(idx, 0).paragraphs[0].add_run(label), 9.5, True, NAVY)
    font(bank.cell(idx, 1).paragraphs[0].add_run(value), 9.5)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(7)
p.paragraph_format.space_after = Pt(2)
font(p.add_run("Срок оплаты"), 10.5, True)
font(p.add_run("  В течение 5 рабочих дней после подписания договора. Счет действителен при неизменности указанных реквизитов."), 10.2)

sign = doc.add_table(rows=1, cols=2)
sign.alignment = WD_TABLE_ALIGNMENT.CENTER
sign.autofit = False
set_table_borders(sign, color="FFFFFF", size="0")
left, right = sign.rows[0].cells
set_width(left, 3.45)
set_width(right, 3.45)
for c in (left, right):
    set_cell_margins(c, top=150, bottom=80)
font(left.paragraphs[0].add_run("Поставщик"), 9.5, True, MUTED)
p = left.add_paragraph()
font(p.add_run("ТОО «ELC ALMATY»"), 10.5, True)
p = left.add_paragraph()
p.paragraph_format.space_before = Pt(10)
font(p.add_run("________________  Цай П. Л."), 10)
font(right.paragraphs[0].add_run("Назначение платежа"), 9.5, True, MUTED)
p = right.add_paragraph()
font(p.add_run("Предоплата по договору № PNK-2026-01. Без НДС."), 10)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(0)
font(fp.add_run("Счет № PNK-С-2026-01   ТОО ELC ALMATY   pllato.kz"), 8.5, False, MUTED)

doc.core_properties.title = "Счет на предоплату № PNK-С-2026-01"
doc.core_properties.subject = "Предоплата 10 процентов по договору PNK-2026-01"
doc.core_properties.author = "ТОО ELC ALMATY"
doc.core_properties.keywords = "ПНК, счет, предоплата, ELC ALMATY"

OUT.parent.mkdir(parents=True, exist_ok=True)
APP_OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
doc.save(APP_OUT)
print(OUT)
print(APP_OUT)
